"""Read a reviewed ontology workbook and export v0.3 package JSON offline.

This module only reads XLSX with openpyxl. It does not create/edit workbooks,
connect to application servers, or import definitions into a live database.
"""
from __future__ import annotations
import argparse
from datetime import date,datetime
import hashlib
import json
import math
from pathlib import Path
import re
import sys
import zipfile
from urllib.parse import quote

from validate_packages import dependency_order,validate_packages

SHEETS={
 '领域包':('packages',['enabled','package_id','version','namespace','label','description']),
 '对象类型':('types',['enabled','package_id','version','type_id','identity_type','label','description']),
 '属性定义':('properties',['enabled','package_id','version','property_id','label','description','applies_to','value_type','unit','enum_values','usage','required','cardinality','minimum','maximum']),
 '关系类型':('predicates',['enabled','package_id','version','predicate_id','label','description','subject_types','object_types','subject_semantic_types','object_semantic_types','dimension','irreflexive','acyclic']),
 '包依赖':('dependencies',['enabled','package_id','version','depends_on_package_id','depends_on_version']),
}
REQUIRED={
 'packages':{'package_id','version','namespace','label'},
 'types':{'package_id','version','type_id','identity_type','label'},
 'properties':{'package_id','version','property_id','label','applies_to','value_type','usage'},
 'predicates':{'package_id','version','predicate_id','subject_types','object_types','dimension'},
 'dependencies':{'package_id','version','depends_on_package_id','depends_on_version'},
}
ARRAY_FIELDS={'applies_to','enum_values','subject_types','object_types','subject_semantic_types','object_semantic_types'}
BOOL_FIELDS={'required','irreflexive','acyclic'}
SUPPORTED_HANDLING={'scalar','foreign_key','array_scalar','enum','identity'}
MAX_ROWS=50000
MAX_COLUMNS=100

class WorkbookError(ValueError):pass

def _where(row):return str(row.get('__sheet','?'))+'!'+str(row.get('__row','?'))

def _error(row,message):raise WorkbookError(_where(row)+': '+message)

def _blank(value):return value is None or value==''

def _string(row,key,required=False):
    value=row.get(key)
    if _blank(value):
        if required:_error(row,'缺少 '+key)
        return None
    if not isinstance(value,str):_error(row,key+' 必须是文本；标识/版本不能用Excel数值或日期代替')
    if value!=value.strip():_error(row,key+' 含首尾空白，请明确修正后再导出')
    if not value.strip():_error(row,key+' 不能为空白')
    return value

def _boolean(row,key,default=False):
    value=row.get(key)
    if _blank(value):return default
    if not isinstance(value,bool):_error(row,key+' 必须是Excel原生 TRUE/FALSE 布尔值')
    return value

def _array(row,key,required=False):
    text=_string(row,key,required)
    if text is None:return None
    try:value=json.loads(text)
    except (ValueError,TypeError):_error(row,key+' 必须为JSON数组文本，如 ["type:Field"]')
    if not isinstance(value,list) or any(not isinstance(x,str) or not x or x!=x.strip() for x in value):_error(row,key+' 必须为非空字符串组成的JSON数组')
    if len(set(value))!=len(value):_error(row,key+' 不允许重复成员')
    if required and not value:_error(row,key+' 不能为空数组')
    return value or None

def _number(row,key):
    value=row.get(key)
    if _blank(value):return None
    if isinstance(value,bool) or not isinstance(value,(int,float)) or not math.isfinite(value):_error(row,key+' 必须是有限Excel数值')
    return value

def _check_handling(row,kind):
    handling=_string(row,'handling')
    if handling and (handling.startswith('deferred') or handling in {'unsupported','excluded','manual_review'}):
        _error(row,'handling='+handling+' 不可启用；保留 enabled=FALSE 或先完成显式建模')
    if kind!='properties':return
    sql=(_string(row,'source_sql_type') or '').lower().strip()
    if handling and handling not in SUPPORTED_HANDLING:_error(row,'未知属性 handling='+handling)
    if re.search(r'\b(jsonb?|geometry|geography|hstore|bytea|xml|record|vector|halfvec|sparsevec|polygon|linestring|point)\b',sql):
        _error(row,'SQL JSON/几何/二进制等复合字段不能默默转换普通属性；请停用并保留原件/关系或显式字段映射')
    is_array='[]' in sql or bool(re.search(r'\barray\b',sql))
    if is_array and (handling!='array_scalar' or row.get('cardinality')!='many'):
        _error(row,'SQL数组必须显式 handling=array_scalar 且 cardinality=many，并定义元素类型')
    if handling=='array_scalar' and row.get('cardinality')!='many':_error(row,'array_scalar 必须为many')
    target=_string(row,'value_type',True)
    # Check obvious incompatible scalar declarations; custom SQL enum/domain names
    # still need an explicit supported target type and are reviewed in the workbook.
    base=sql.replace('[]','').strip()
    if re.match(r'^(smallint|integer|int|int2|int4|int8|bigint|smallserial|serial|bigserial)\b',base) and target!='integer':_error(row,'SQL整数应映射integer，不能隐式转为其他值类型')
    if re.match(r'^(numeric|decimal|real|double precision|float[48]?)\b',base) and target!='decimal':_error(row,'SQL数值应映射decimal')
    if re.match(r'^(bool|boolean)\b',base) and target!='boolean':_error(row,'SQL布尔应映射boolean')
    if re.match(r'^(timestamp|timestamptz)\b',base) and target!='datetime':_error(row,'SQL timestamp应映射datetime；时区仍须在接入映射中明确')
    if target=='datetime' and re.match(r'^(date|time)\b',base):_error(row,'只有日期/时刻的SQL值不能自动冒充完整带时区datetime')


def _bounded_sheet_rows(ws,sheet):
    """Iterate actual rows even when a producer omitted/understated dimension."""
    if ((ws.max_row is not None and ws.max_row>MAX_ROWS+2)
            or (ws.max_column is not None and ws.max_column>MAX_COLUMNS)):
        raise WorkbookError(sheet+' 超过读取预算')
    # This only resets reader metadata; it never edits the XLSX. Without it an
    # understated producer dimension could silently hide later rows/columns.
    ws.reset_dimensions()
    for number,cells in enumerate(ws.iter_rows(),1):
        if number>MAX_ROWS+2 or len(cells)>MAX_COLUMNS:
            raise WorkbookError(sheet+' 实际行列超过读取预算')
        yield cells


def read_workbook(path):
    """Read original cell values; formulas are rejected instead of trusting caches."""
    from openpyxl import load_workbook
    path=Path(path)
    if path.suffix.lower()!='.xlsx':raise WorkbookError('仅支持无宏.xlsx文件')
    if path.stat().st_size>64*1024*1024:raise WorkbookError('工作簿超过64MiB读取上限')
    with zipfile.ZipFile(path) as archive:
        members=archive.infolist()
        if len(members)>2000 or sum(x.file_size for x in members)>256*1024*1024:raise WorkbookError('工作簿展开预算超限')
        if any('vbaproject' in x.filename.lower() for x in members):raise WorkbookError('不接受宏工作簿')
    book=load_workbook(path,read_only=True,data_only=False,keep_links=False)
    result={}
    try:
        for sheet,(kind,fields) in SHEETS.items():
            if sheet not in book.sheetnames:raise WorkbookError('缺少工作表 '+sheet)
            ws=book[sheet]
            iterator=_bounded_sheet_rows(ws,sheet)
            head=next(iterator,())
            names=[]
            for cell in head:
                value=cell.value
                if value is None:names.append(None);continue
                if cell.data_type in {'f','e'} or not isinstance(value,str) or value!=value.strip():raise WorkbookError(sheet+'!1 机器表头不合法')
                names.append(value)
            nonempty=[x for x in names if x]
            if len(nonempty)!=len(set(nonempty)):raise WorkbookError(sheet+'!1 重复表头')
            missing=set(fields)-set(nonempty)
            if missing:raise WorkbookError(sheet+'!1 缺少机器列 '+','.join(sorted(missing)))
            # Row two is the human-readable Chinese caption row, not data.
            if not next(iterator,None):raise WorkbookError(sheet+' 缺少第2行中文说明表头')
            rows=[]
            for number,cells in enumerate(iterator,3):
                if all(cell.value is None for cell in cells):continue
                row={'__sheet':sheet,'__row':number}
                if any(cell.value is not None for cell in cells[len(names):]):_error(row,'机器表头之外的列包含数据')
                for name,cell in zip(names,cells):
                    if not name:
                        if cell.value is not None:_error(row,'无机器表头列包含数据')
                        continue
                    # Formula and Excel error values in import/control columns fail
                    # closed. Provenance-only explanatory columns never execute.
                    if name in set(fields)|{'handling','source_sql_type'} and cell.data_type in {'f','e'}:_error(row,name+' 不可使用公式或Excel错误值')
                    row[name]=cell.value
                rows.append(row)
            result[kind]=rows
    finally:book.close()
    return result


def rows_to_packages(tables):
    packages={};locations={};skipped=[];active={}
    for kind in REQUIRED:
        active[kind]=[]
        for index,raw in enumerate(tables.get(kind,[]),3):
            row=dict(raw);row.setdefault('__sheet',kind);row.setdefault('__row',index)
            if not isinstance(row.get('enabled'),bool):_error(row,'enabled 必须为Excel原生 TRUE/FALSE')
            if not row['enabled']:
                skipped.append({'sheet':row['__sheet'],'row':row['__row'],'package_id':row.get('package_id'),'id':row.get('property_id') or row.get('predicate_id') or row.get('type_id') or row.get('package_id'),'handling':row.get('handling'),'reason':'enabled=FALSE'})
                continue
            _check_handling(row,kind)
            active[kind].append(row)
    for row in active['packages']:
        body={key:_string(row,key,key in REQUIRED['packages']) for key in ['package_id','version','namespace','label','description']}
        body={k:v for k,v in body.items() if v is not None};key=(body['package_id'],body['version'])
        if key in packages:_error(row,'同一本体包版本重复，首次见 '+locations[key])
        body.update(types=[],properties=[],predicates=[],dependencies=[])
        packages[key]=body;locations[key]=_where(row)
    if not packages:raise WorkbookError('没有启用的领域包')
    seen={kind:set() for kind in ['types','properties','predicates','dependencies']}
    for kind in ['types','properties','predicates','dependencies']:
        allowed=next(fields for k,fields in SHEETS.values() if k==kind)
        for row in active[kind]:
            key=(_string(row,'package_id',True),_string(row,'version',True))
            if key not in packages:_error(row,'行引用的领域包未启用或不存在: '+repr(key))
            body={}
            for column in allowed:
                if column in {'enabled','package_id','version'}:continue
                required=column in REQUIRED[kind]
                if column in ARRAY_FIELDS:value=_array(row,column,required)
                elif column in BOOL_FIELDS:value=_boolean(row,column)
                elif column in {'minimum','maximum'}:value=_number(row,column)
                else:value=_string(row,column,required)
                if value is not None:body[column]=value
            if kind=='dependencies':
                body={'package_id':body['depends_on_package_id'],'version':body['depends_on_version']}
                identity=key+(body['package_id'],body['version'])
            else:
                idkey={'types':'type_id','properties':'property_id','predicates':'predicate_id'}[kind]
                identity=key+(body[idkey],)
            if identity in seen[kind]:_error(row,'重复定义标识: '+repr(identity))
            seen[kind].add(identity);packages[key][kind].append(body)
    ordered=dependency_order(list(packages.values()))
    # Semantic checks (class membership, units, enums, endpoint identity) are run
    # through the current service in an isolated DB, never guessed from SQL text.
    validation=validate_packages(ordered)
    return ordered,{'format':'ontology-workbook/1','status':'passed','skipped_rows':skipped,'validation':validation,
        'limitations':['仅导出本体结构，未生成现场业务对象或导入生产数据库','SQLite隔离验证不替代PostgreSQL、真实映射或现场验收','日期、JSON、空间、数组和跨包语义限制需按工作簿说明显式处理']}


def write_output(packages,report,workbook,output):
    output=Path(output);output.mkdir(parents=True,exist_ok=True)
    payloads=[]
    for index,package in enumerate(packages,1):
        name=f'{index:03d}_'+quote(package['package_id'],safe='._-')[:100]+'_'+quote(package['version'],safe='._-')+'.json'
        data=(json.dumps(package,ensure_ascii=False,indent=2,allow_nan=False)+'\n').encode('utf-8')
        (output/name).write_bytes(data)
        payloads.append({'package_id':package['package_id'],'version':package['version'],'file':name,'sha256':hashlib.sha256(data).hexdigest(),
                         'method':'POST','endpoint':'/api/v1/ontology/packages'})
    manifest={'format':'ontology-package-export/1','source_workbook':Path(workbook).name,'workbook_sha256':hashlib.sha256(Path(workbook).read_bytes()).hexdigest(),
              'import_order':'packages数组顺序；精确版本依赖在前','automatic_import':False,'packages':payloads}
    (output/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    (output/'validation-report.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    return manifest


def main(argv=None):
    parser=argparse.ArgumentParser(description='只读Excel并离线导出本体包JSON，不访问现有数据库或网络')
    parser.add_argument('workbook',type=Path)
    parser.add_argument('--output-dir',type=Path,help='默认在工作簿旁创建<名称>_packages目录')
    parser.add_argument('--check-only',action='store_true',help='只在隔离SQLite校验，不输出JSON文件')
    args=parser.parse_args(argv)
    try:
        packages,report=rows_to_packages(read_workbook(args.workbook))
        if not args.check_only:
            destination=args.output_dir or args.workbook.with_name(args.workbook.stem+'_packages')
            manifest=write_output(packages,report,args.workbook,destination)
            print(json.dumps({'status':'passed','output_dir':str(destination.resolve()),'files':len(manifest['packages']),'counts':report['validation']['definition_counts'],'skipped_rows':len(report['skipped_rows']),'database_imported':False},ensure_ascii=False,indent=2))
        else:print(json.dumps(report,ensure_ascii=False,indent=2))
        return 0
    except Exception as exc:
        print(json.dumps({'status':'failed','error':str(exc),'database_imported':False},ensure_ascii=False),file=sys.stderr);return 2

if __name__=='__main__':raise SystemExit(main())
