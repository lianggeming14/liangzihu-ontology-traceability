from pathlib import Path
import json,hashlib,zipfile,xml.etree.ElementTree as ET
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[2];OUT=ROOT/'outputs/liangzihu-ontology-20261015'
p=OUT/'梁子湖项目_本体与关系全量结构导出.xlsx'
w=load_workbook(p,read_only=True,data_only=False,keep_links=False)
d=json.loads((OUT/'workbook_data.json').read_text(encoding='utf8'))
counts={};error_cells=[];formulas=[]
for s in w:
 s.reset_dimensions()
 for row in s.iter_rows():
  for c in row:
   if c.data_type=='e':error_cells.append(s.title+'!'+c.coordinate)
   if c.data_type=='f':formulas.append(s.title+'!'+c.coordinate)
for spec in d['sheets']:
 s=w[spec['name']];it=s.iter_rows(values_only=True);head=next(it);next(it);actual=[dict(zip(head,r)) for r in it if any(v is not None for v in r)]
 assert len(actual)==len(spec['rows']),(s.title,len(actual),len(spec['rows']))
 for a,b in zip(actual,spec['rows']):
  for k,v in b.items():
   if v=='':v=None
   assert a.get(k)==v,(s.title,k,a.get(k),v)
 counts[s.title]=len(actual)
for name in ['试录对象','试录属性','试录关系']:
 s=w[name];rows=list(s.iter_rows())
 assert all(c.value is None for row in rows[2:] for c in row),name
 assert all(c.number_format=='@' for row in rows[2:] for c in row if c.has_style),name
w.close()
assert not error_cells and not formulas
m=json.loads((OUT/'packages/manifest.json').read_text(encoding='utf8'))
assert hashlib.sha256(p.read_bytes()).hexdigest()==m['workbook_sha256']
ns={'x':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
panes=[];validations=0
with zipfile.ZipFile(p) as z:
 assert z.testzip() is None
 for name in z.namelist():
  if name.startswith('xl/worksheets/sheet') and name.endswith('.xml'):
   x=ET.fromstring(z.read(name));panes.append(x.find('x:sheetViews/x:sheetView/x:pane',ns).attrib)
   validations+=len(x.findall('x:dataValidations/x:dataValidation',ns))
assert all(float(q.get('ySplit',0))>=2 for q in panes),panes
r={'status':'passed','sheets':len(panes),'source_row_counts':counts,'excel_error_cells':error_cells,'formula_cells':formulas,'blank_instance_templates':True,'text_format_preserves_input_tokens':True,'data_validations':validations,'frozen_headers':True,'package_manifest_matches_workbook':True,'workbook_sha256':m['workbook_sha256'],'database_imported':False}
(OUT/'final-workbook-verification.json').write_text(json.dumps(r,ensure_ascii=False,indent=2)+'\n',encoding='utf8')
print(json.dumps(r,ensure_ascii=False,indent=2))
