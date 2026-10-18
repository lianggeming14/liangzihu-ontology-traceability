"""Build editable Chinese operation/report DOCX and deterministic layout QA previews.

The PDF is a measured layout preview, NOT evidence of a native Word render.
Run the packaged render_docx.py separately and record its actual result.
"""
from pathlib import Path
import argparse
import hashlib
import json
import re
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas
from reportlab.lib.colors import HexColor

W, H = 595.276, 841.89
MARGIN, TOP, BOTTOM = 55.0, 48.0, 57.0
WIDTH, HEIGHT = W - 2 * MARGIN, H - TOP - BOTTOM - 15


def xml(tag, **attrs):
    element = OxmlElement(tag)
    for key, value in attrs.items(): element.set(qn(key), str(value))
    return element


def font(run, size=10.5, bold=False):
    name = '微软雅黑' if bold else '宋体'
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    run._element.get_or_add_rPr().append(xml('w:rFonts', **{'w:ascii':name, 'w:hAnsi':name, 'w:eastAsia':name}))


def wrap(value, width=WIDTH, size=10.5, bold=False):
    face = 'DocHei' if bold else 'DocSong'
    output = []
    for paragraph in str(value).split('\n'):
        current = ''
        for token in re.findall(r'[A-Za-z0-9_./:#@+<>%=-]+|.', paragraph):
            chunks = list(token) if pdfmetrics.stringWidth(token, face, size) > width else [token]
            for piece in chunks:
                if current and pdfmetrics.stringWidth(current+piece, face, size) > width:
                    if piece in '，。；：？！、）》】' and len(current)>1:
                        output.append(current[:-1]); current=current[-1]+piece
                    else:
                        output.append(current.rstrip()); current=piece.lstrip()
                else: current += piece
        output.append(current)
    return output


class Layout:
    def __init__(self): self.pages=[[]]; self.used=0
    def new(self): self.pages.append([]); self.used=0
    def add(self, item, height):
        # Fail visibly rather than creating a preview whose body crosses the
        # reserved footer area. The 15 pt reserve also covers Word's small
        # paragraph/table placement differences; native rendering is separate.
        if height < 0 or self.used + height > HEIGHT + 0.01:
            raise ValueError(f"Page {len(self.pages)} exceeds body budget: {self.used + height:.2f} > {HEIGHT:.2f} pt")
        item.update(top=self.used, height=height)
        self.pages[-1].append(item); self.used += height
    def text(self, text, kind='p'):
        size, lead, before, after = {
            'title':(22,31,8,14), 'subtitle':(12,19,0,12), 'h1':(14,22,14,8),
            'h2':(11.5,19,9,5), 'code':(9,13.5,0,2), 'source':(9,14,0,4)
        }.get(kind,(10.5,17,0,7))
        bold=kind in ('title','h1','h2')
        lines=wrap(text, size=size, bold=bold)
        reserve=55 if kind in ('h1','h2') else 0
        if self.used+before+lead*min(len(lines),3)+after+reserve>HEIGHT and self.used:
            self.new(); before=0
        while lines:
            count=int((HEIGHT-self.used-before-after)/lead)
            if count < 1:
                if not self.used:
                    raise ValueError('Text line exceeds the available body height')
                self.new(); before=0
                continue
            if count<len(lines) and len(lines)-count==1 and count>2: count-=1
            part, lines=lines[:count],lines[count:]
            self.add(dict(kind=kind,lines=part,size=size,lead=lead,bold=bold,before=before,after=after),before+len(part)*lead+after)
            if lines: self.new(); before=0
    def table(self, specification):
        cols=specification['headers']; n=len(cols)
        ratios=specification.get('widths',[100/n]*n)
        widths=[WIDTH*r/sum(ratios) for r in ratios]
        def row(values, head=False):
            lines=[wrap(v,w-10,9.5,head) for v,w in zip(values,widths)]
            return dict(cells=lines,height=max(map(len,lines))*14+10,head=head)
        header=row(cols,True); rows=[row(v) for v in specification['rows']]
        if not rows:return
        if self.used+header['height']+rows[0]['height']+7>HEIGHT:self.new()
        part=[header];height=header['height']
        for rr in rows:
            if header['height']+rr['height']>HEIGHT:raise ValueError('Table row too tall; write this content as prose')
            if self.used+height+rr['height']+7>HEIGHT:
                self.add(dict(kind='table',widths=widths,rows=part),height+7)
                self.new();part=[header];height=header['height']
            part.append(rr);height+=rr['height']
        self.add(dict(kind='table',widths=widths,rows=part),height+7)


def render_docx(data, layout, output):
    doc=Document(); sec=doc.sections[0]
    sec.page_width=Pt(W);sec.page_height=Pt(H)
    sec.left_margin=sec.right_margin=Pt(MARGIN)
    sec.top_margin=Pt(TOP);sec.bottom_margin=Pt(BOTTOM)
    sec.footer_distance=Pt(24)
    for name in ('Normal','Title','Subtitle','Heading 1','Heading 2'):
        style=doc.styles[name];style.font.color.rgb=RGBColor(0,0,0)
        style.font.name='宋体';style.font.size=Pt(10.5)
        pf=style.paragraph_format;pf.space_before=pf.space_after=Pt(0)
        pf.keep_with_next=False;pf.widow_control=False;pf.page_break_before=False
        style.element.get_or_add_rPr().append(xml('w:rFonts',**{'w:ascii':'宋体','w:hAnsi':'宋体','w:eastAsia':'宋体'}))
        for element in list(style.element.xpath('.//w:pBdr')): element.getparent().remove(element)
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    font(footer.add_run('梁子湖农业数据工程    '),8.5)
    run=footer.add_run();font(run,8.5)
    run._r.append(xml('w:fldChar',**{'w:fldCharType':'begin'}))
    instruction=xml('w:instrText');instruction.text=' PAGE ';run._r.append(instruction)
    run._r.append(xml('w:fldChar',**{'w:fldCharType':'end'}))
    for pi,page in enumerate(layout.pages):
        if pi:
            para=doc.add_paragraph();para.paragraph_format.page_break_before=True
            para.paragraph_format.line_spacing=Pt(1)
            font(para.add_run(''),1)
        for item in page:
            if item['kind']=='table':
                table=doc.add_table(rows=0,cols=len(item['widths']));table.autofit=False
                table.alignment=WD_TABLE_ALIGNMENT.CENTER
                for col,width in zip(table.columns,item['widths']):col.width=Pt(width)
                borders=xml('w:tblBorders')
                for edge in ('top','left','bottom','right','insideH','insideV'):
                    borders.append(xml('w:'+edge,**{'w:val':'single','w:sz':'4','w:color':'D9D9D9'}))
                table._tbl.tblPr.append(borders)
                margins=xml('w:tblCellMar')
                for side in ('top','left','bottom','right'):margins.append(xml('w:'+side,**{'w:w':'100','w:type':'dxa'}))
                table._tbl.tblPr.append(margins)
                for ri,row in enumerate(item['rows']):
                    result=table.add_row();result._tr.get_or_add_trPr().append(xml('w:cantSplit'))
                    if row['head']:result._tr.get_or_add_trPr().append(xml('w:tblHeader'))
                    for cell,width,lines in zip(result.cells,item['widths'],row['cells']):
                        cell.width=Pt(width);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                        cell._tc.get_or_add_tcPr().append(xml('w:shd',**{'w:fill':'DDEAF2' if row['head'] else ('F4F7F9' if ri%2==0 else 'FFFFFF')}))
                        para=cell.paragraphs[0];para.paragraph_format.line_spacing=Pt(14)
                        para.paragraph_format.keep_with_next=False
                        for i,line in enumerate(lines):
                            if i:para.add_run().add_break()
                            font(para.add_run(line),9.5,row['head'])
                para=doc.add_paragraph();para.paragraph_format.line_spacing=Pt(7)
            else:
                para=doc.add_paragraph(style={'title':'Title','subtitle':'Subtitle','h1':'Heading 1','h2':'Heading 2'}.get(item['kind'],'Normal'))
                pf=para.paragraph_format;pf.line_spacing=Pt(item['lead']);pf.space_before=Pt(item['before']);pf.space_after=Pt(item['after'])
                pf.keep_together=True;pf.keep_with_next=False
                for i,line in enumerate(item['lines']):
                    if i:para.add_run().add_break()
                    font(para.add_run(line),item['size'],item['bold'])
    for para in list(doc.paragraphs)+[p for t in doc.tables for r in t.rows for c in r.cells for p in c.paragraphs]:
        props=para._p.get_or_add_pPr()
        for tag in ('snapToGrid','autoSpaceDE','autoSpaceDN'):props.append(xml('w:'+tag,**{'w:val':'0'}))
    doc.core_properties.title=data['title'];doc.core_properties.author='';doc.core_properties.subject=data.get('subtitle','')
    output.parent.mkdir(parents=True,exist_ok=True);doc.save(output)


def render_preview(data,layout,output):
    pdf=canvas.Canvas(str(output),pagesize=(W,H));pdf.setTitle(data['title']);pdf.setAuthor('')
    for pi,page in enumerate(layout.pages,1):
        for item in page:
            y=H-TOP-item['top']
            if item['kind']=='table':
                for ri,row in enumerate(item['rows']):
                    x=MARGIN
                    for width,lines in zip(item['widths'],row['cells']):
                        pdf.setFillColor(HexColor('#DDEAF2' if row['head'] else ('#F4F7F9' if ri%2==0 else '#FFFFFF')))
                        pdf.setStrokeColor(HexColor('#D9D9D9'));pdf.setLineWidth(.5)
                        pdf.rect(x,y-row['height'],width,row['height'],fill=1,stroke=1)
                        pdf.setFillColor(HexColor('#000000'));pdf.setFont('DocHei' if row['head'] else 'DocSong',9.5)
                        yy=y-5-9.5
                        for line in lines:pdf.drawString(x+5,yy,line);yy-=14
                        x+=width
                    y-=row['height']
            else:
                pdf.setFillColor(HexColor('#000000'));pdf.setFont('DocHei' if item['bold'] else 'DocSong',item['size'])
                y-=item['before']+item['size']
                for line in item['lines']:pdf.drawString(MARGIN,y,line);y-=item['lead']
        pdf.setFont('DocSong',8.5);pdf.drawCentredString(W/2,27,'梁子湖农业数据工程    '+str(pi))
        pdf.showPage()
    pdf.save()


def main():
    parser=argparse.ArgumentParser();parser.add_argument('input');parser.add_argument('output');parser.add_argument('--qa-dir',required=True)
    args=parser.parse_args();data=json.loads(Path(args.input).read_text(encoding='utf-8'))
    pdfmetrics.registerFont(TTFont('DocSong','C:/Windows/Fonts/simsun.ttc',subfontIndex=0))
    pdfmetrics.registerFont(TTFont('DocHei','C:/Windows/Fonts/msyh.ttc',subfontIndex=0))
    layout=Layout();layout.text(data['title'],'title');layout.text(data.get('subtitle','')+'    2026 年 10 月 19 日','subtitle')
    for text in data.get('intro',[]):layout.text(text)
    for section in data['sections']:
        if section.get('page_break_before') and layout.used:
            layout.new()
        layout.text(section['title'],'h1')
        for text in section.get('paragraphs',[]):layout.text(text)
        for table in section.get('tables',[]):layout.table(table)
        for code in section.get('code',[]):layout.text(code,'code')
        for source in section.get('sources',[]):layout.text(source if isinstance(source,str) else json.dumps(source,ensure_ascii=False),'source')
    references=data.get('references',[])
    if references:
        layout.text('附录 代码与结构证据索引','h1')
        for ref in references:
            if isinstance(ref,str):layout.text(ref,'source')
            else:
                # Preserve source notes but avoid printing a JSON dictionary as prose.
                title=ref.get('title') or ref.get('path') or ref.get('id','资料')
                layout.text(str(title),'h2')
                for key in ('id','kind','path','line','lines','evidence_level','claim','note','url'):
                    if ref.get(key):layout.text(str(ref[key]),'source')
    output=Path(args.output);qa=Path(args.qa_dir);qa.mkdir(parents=True,exist_ok=True)
    render_docx(data,layout,output);render_preview(data,layout,qa/'layout-preview.pdf')
    report={'docx':output.name,'sha256':hashlib.sha256(output.read_bytes()).hexdigest(),'planned_pages':len(layout.pages),'native_render_verified':False,'preview_method':'Measured explicit line and page layout using ReportLab; not a native Word render','body_budget_pt':HEIGHT,'minimum_reserved_body_bottom_pt':H-TOP-HEIGHT,'footer_baseline_pt':27,'pages':layout.pages}
    (qa/'layout-report.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({k:v for k,v in report.items() if k!='pages'},ensure_ascii=False))


if __name__=='__main__':main()
