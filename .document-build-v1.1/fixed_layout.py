import json, math, re
from pathlib import Path
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas
from reportlab.lib.colors import HexColor
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
BASE=Path(__file__).resolve().parent; ROOT=BASE.parent
D=json.loads((BASE/'plan-content.json').read_text(encoding='utf-8'))
OUT=ROOT/'交付文档'/'梁子湖智慧农业本体库与溯源系统总体方案_v1.1.docx'
pdfmetrics.registerFont(TTFont('Song','C:/Windows/Fonts/simsun.ttc',subfontIndex=0))
pdfmetrics.registerFont(TTFont('Hei','C:/Windows/Fonts/msyh.ttc',subfontIndex=0))
W,H=595.276,841.89; LM,RM,TM,BM=59.53,59.53,56.70,56.70; CW=W-LM-RM; LIMIT=H-TM-BM-26
PAGES=[[],[]]; used=0.; toc=[]

def measure(s,font,size):return pdfmetrics.stringWidth(s,font,size)
def wrap(text,width,font='Song',size=11):
    lines=[]
    for raw in str(text).split('\n'):
        if not raw:lines.append('');continue
        cur=''
        tokens=re.findall(r'[A-Za-z0-9_./:#@+<>-]+|.',raw)
        expanded=[]
        for token in tokens:
            expanded.extend(list(token) if measure(token,font,size)>width else [token])
        for char in expanded:
            if measure(cur+char,font,size)>width and cur:
                # Keep punctuation with preceding CJK text when practical.
                if char in '，。；：？！、）》】' and len(cur)>1:
                    lines.append(cur[:-1]);cur=cur[-1]+char
                else:lines.append(cur.rstrip());cur=char.lstrip()
            else:cur+=char
        if cur:lines.append(cur)
    return lines

def newpage():
    global used
    PAGES.append([]);used=0.
def add(item,height):
    global used
    item['height']=height;item['top']=used;PAGES[-1].append(item);used+=height

def text_item(text,kind='p',size=11,font='Song',leading=17,after=6,before=0,indent=0,chapter_id=None):
    global used
    lines=wrap(text,CW-indent,font,size)
    ishead=kind in ['h1','h2','h3']
    need=before+len(lines)*leading+after
    reserve=85 if ishead else (15 if kind=='ref' else 0)
    if used+need+reserve>LIMIT and used>0:
        if not ishead and len(lines)>5 and LIMIT-used>leading*3+before+after:
            n=int((LIMIT-used-before-after)/leading)
            if len(lines)-n<2:n=len(lines)-2
            add({'kind':kind,'lines':lines[:n],'text':text,'size':size,'font':font,'leading':leading,'before':before,'after':after,'indent':indent},before+n*leading+after)
            newpage();lines=lines[n:];before=0;need=len(lines)*leading+after
        else:newpage();before=0;need=len(lines)*leading+after
    item={'kind':kind,'lines':lines,'text':text,'size':size,'font':font,'leading':leading,'before':before,'after':after,'indent':indent}
    if chapter_id:item['chapter_id']=chapter_id
    add(item,need)

def table(b):
    global used
    widths=[CW*w/100 for w in b.get('widths',[100/len(b['headers'])]*len(b['headers']))]
    size=9.5;lead=13.2;pad=5
    def row(vals,ishead=False):
        ls=[wrap(v,w-2*pad,'Hei' if ishead else 'Song',size) for v,w in zip(vals,widths)]
        return {'cells':ls,'h':max(map(len,ls))*lead+2*pad,'ishead':ishead}
    header=row(b['headers'],True);rows=[row(r) for r in b['rows']]
    if used+header['h']+rows[0]['h']+4>LIMIT:newpage()
    part=[header];ph=header['h']
    for rr in rows:
        if used+ph+rr['h']+4>LIMIT:
            add({'kind':'table','widths':widths,'rows':part,'size':size,'leading':lead,'pad':pad},ph+6)
            newpage();part=[header];ph=header['h']
        part.append(rr);ph+=rr['h']
    add({'kind':'table','widths':widths,'rows':part,'size':size,'leading':lead,'pad':pad},ph+6)

for ci,c in enumerate(D['chapters'],1):
    if ci==1 or used>LIMIT-170:newpage()
    toc.append({'title':c['title'],'page':len(PAGES),'anchor':'chapter_'+str(ci)})
    text_item(c['title'],'h1',16,'Hei',23,12,20 if used else 0,chapter_id='chapter_'+str(ci))
    for b in c['blocks']:
        k=b['type']
        if k=='p':text_item(b['text'])
        elif k=='h2':text_item(b['text'],'h2',12.5,'Hei',19,7,9)
        elif k=='h3':text_item(b['text'],'h3',11.5,'Hei',18,6,7)
        elif k=='bullets':
            for s in b['items']:text_item('• '+s,'bullet',11,'Song',17,5,0,0)
        elif k=='table':table(b)
        elif k=='code':
            for s in b['text'].splitlines():text_item(s or ' ','code',9,'Song',12,0)
            text_item(' ','spacer',3,'Song',3,4)
        elif k=='reference':
            s='['+b['id']+'] '+b['title']+'。'+b['note']
            text_item(s,'ref',10,'Song',15,3)
            if b['url']:
                text_item('原始资料','link',10,'Song',14,1);PAGES[-1][-1]['url']=b['url']
        elif k=='figure':
            ratio=1080/1600 if b['name']=='architecture' else 880/1600
            ih=CW*ratio; need=ih+25
            if used+need>LIMIT:newpage()
            add({'kind':'figure','path':str(BASE/(b['name']+'.png')),'image_h':ih,'caption':b['caption']},need)
# Cover and static linked table of contents.
PAGES[0]=[
 {'kind':'cover_title','top':60,'lines':wrap(D['title'],CW,'Hei',24),'size':24,'font':'Hei','leading':36},
 {'kind':'cover_sub','top':160,'lines':wrap(D['subtitle'],CW,'Hei',14),'size':14,'font':'Hei','leading':22},
 {'kind':'cover_meta','top':203,'lines':['版本 '+D['version']+'     '+D['date']],'size':11,'font':'Song','leading':18},
]
def coverpara(top,text,kind='p',size=11,font='Song',leading=18):PAGES[0].append({'kind':kind,'top':top,'lines':wrap(text,CW,font,size),'size':size,'font':font,'leading':leading})
coverpara(260,'建设重点','h2',13,'Hei',20)
coverpara(293,'构建从信源接收、语义对齐、事实与证据管理，到代码层推理输入的完整数据基础。统一五大语义重点，以虾稻为试点并支持多品类扩展，将证据血缘与产品物料谱系连接。')
coverpara(385,'建议保留 PostgreSQL 作为权威存储，在现有梁子湖农场平台上新增语义、证据及推理接口模块。属性图与 RDF 作为可重建和可交换的表示。先验收数据库与供数接口，再建设内部溯源工作台、审核发布和消费者查询。')
coverpara(486,'阅读使用','h2',13,'Hei',20)
coverpara(519,'第一至三章说明目标、架构和平台集成；第四至六章定义语义、知识图及标准交换；第七至十一章说明数据处理、冲突、溯源、存储与推理接口；第十二至十三章给出实施与验收；第十四至十六章明确领域扩展、供数接口、产品溯源和 Kimi 执行门槛。附录提供契约与参考资料。')
PAGES[1]=[{'kind':'toc_title','top':0,'lines':['目录'],'size':22,'font':'Hei','leading':30}]
for i,x in enumerate(toc):PAGES[1].append({'kind':'toc','top':55+i*31,'lines':[x['title']],'size':11,'font':'Song','leading':20,**x})
# Render the fixed pagination preview with exactly the wrapped strings used in DOCX.
PDF=BASE/'fixed-layout-preview.pdf';cv=canvas.Canvas(str(PDF),pagesize=(W,H));cv.setTitle(D['title']);cv.setAuthor('')
for pi,page in enumerate(PAGES,1):
 for item in page:
    kind=item['kind'];top=TM+item['top']
    if kind=='table':
        y=H-top
        for ri,rr in enumerate(item['rows']):
            x=LM;rh=rr['h'];cv.setFillColor(HexColor('#E7EEF4' if rr['ishead'] else ('#F7F9FB' if ri%2==0 else '#FFFFFF')));cv.setStrokeColor(HexColor('#D9D9D9'));cv.setLineWidth(.5)
            for width,ls in zip(item['widths'],rr['cells']):
                cv.rect(x,y-rh,width,rh,fill=1,stroke=1);font='Hei' if rr['ishead'] else 'Song';cv.setFont(font,item['size']);cv.setFillColor(HexColor('#000000'))
                yy=y-item['pad']-item['size']
                for line in ls:cv.drawString(x+item['pad'],yy,line);yy-=item['leading']
                x+=width;cv.setFillColor(HexColor('#E7EEF4' if rr['ishead'] else ('#F7F9FB' if ri%2==0 else '#FFFFFF')))
            y-=rh
    elif kind=='figure':
        cv.drawImage(item['path'],LM,H-top-item['image_h'],width=CW,height=item['image_h']);cv.setFont('Song',10);cv.setFillColor(HexColor('#000000'));cv.drawCentredString(W/2,H-top-item['image_h']-17,item['caption'])
    else:
        if item.get('chapter_id'):cv.bookmarkPage(item['chapter_id']);cv.addOutlineEntry(item['text'],item['chapter_id'],0)
        cv.setFont(item['font'],item['size']);cv.setFillColor(HexColor('#1F4E79' if kind=='link' else '#000000'))
        y=H-top-item.get('before',0)-item['size']
        for line in item['lines']:cv.drawString(LM+item.get('indent',0),y,line);y-=item['leading']
        if kind=='toc':
            cv.drawRightString(W-RM,H-top-item['size'],str(item['page']));cv.linkRect('',item['anchor'],(LM,H-top-22,W-RM,H-top+3),relative=0,thickness=0)
        if kind=='link':cv.linkURL(item['url'],(LM,H-top-15,LM+55,H-top+2),relative=0)
 if pi>1:
    cv.setFont('Song',8.5);cv.setFillColor(HexColor('#333333'));cv.drawCentredString(W/2,28,'梁子湖智慧农业本体论数据库    '+str(pi)+' / '+str(len(PAGES)))
 cv.showPage()
cv.save()
# Build editable DOCX using manual line breaks and identical explicit page boundaries.
def e(tag,**attrs):
    x=OxmlElement(tag)
    for k,v in attrs.items():x.set(qn(k),str(v))
    return x

def runfont(run,font,size):
    fn='微软雅黑' if font=='Hei' else '宋体';run.font.name=fn;run.font.size=Pt(size);run.font.color.rgb=RGBColor(0,0,0)
    rp=run._element.get_or_add_rPr();rp.append(e('w:rFonts',**{'w:ascii':fn,'w:hAnsi':fn,'w:eastAsia':fn}))

doc=Document();sec=doc.sections[0];sec.page_width=Pt(W);sec.page_height=Pt(H);sec.top_margin=Pt(TM);sec.bottom_margin=Pt(BM);sec.left_margin=Pt(LM);sec.right_margin=Pt(RM);sec.footer_distance=Pt(23);sec.different_first_page_header_footer=True
for sn in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
    st=doc.styles[sn];st.font.color.rgb=RGBColor(0,0,0);st.font.name='宋体';st.font.size=Pt(11)
    st.paragraph_format.space_before=Pt(0);st.paragraph_format.space_after=Pt(0);st.paragraph_format.line_spacing=Pt(17);st.paragraph_format.keep_with_next=False;st.paragraph_format.widow_control=False
    st.paragraph_format.page_break_before=False
    st.element.get_or_add_rPr().append(e('w:rFonts',**{'w:ascii':'宋体','w:hAnsi':'宋体','w:eastAsia':'宋体'}))
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
rr=footer.add_run('梁子湖智慧农业本体论数据库    ');runfont(rr,'Song',8.5)
rr=footer.add_run();rr._r.append(e('w:fldChar',**{'w:fldCharType':'begin'}));xx=e('w:instrText');xx.text=' PAGE ';rr._r.append(xx);rr._r.append(e('w:fldChar',**{'w:fldCharType':'end'}));runfont(rr,'Song',8.5)
rr=footer.add_run(' / '+str(len(PAGES)));runfont(rr,'Song',8.5)
bookmark=1

def paragraph(item,style=None,space_before=None):
    global bookmark
    pp=doc.add_paragraph(style=style);fmt=pp.paragraph_format
    fmt.line_spacing=Pt(item['leading']);fmt.space_before=Pt(item.get('before',0) if space_before is None else space_before);fmt.space_after=Pt(item.get('after',0));fmt.keep_together=True;fmt.keep_with_next=False;fmt.widow_control=False
    pr=pp._p.get_or_add_pPr();pr.append(e('w:snapToGrid',**{'w:val':'0'}));pr.append(e('w:contextualSpacing',**{'w:val':'0'}))
    pp.paragraph_format.left_indent=Pt(item.get('indent',0))
    if item.get('chapter_id'):
        pp._p.append(e('w:bookmarkStart',**{'w:id':bookmark,'w:name':item['chapter_id']}));pp._p.append(e('w:bookmarkEnd',**{'w:id':bookmark}));bookmark+=1
    for j,line in enumerate(item['lines']):
        if j:pp.add_run().add_break()
        rr=pp.add_run(line);runfont(rr,item['font'],item['size'])
    return pp

for pi,page in enumerate(PAGES):
    if pi:
        pb=doc.add_paragraph();pb.paragraph_format.line_spacing=Pt(1);pb.paragraph_format.space_before=Pt(0);pb.paragraph_format.space_after=Pt(0);pb.add_run().add_break(WD_BREAK.PAGE)
    prev_bottom=0
    for item in page:
        kind=item['kind']
        if pi<2:
            gap=max(0,item['top']-prev_bottom)
            if kind=='toc':
                pp=paragraph(item,space_before=gap);pp.paragraph_format.tab_stops.add_tab_stop(Pt(CW),WD_ALIGN_PARAGRAPH.RIGHT)
                rr=pp.add_run('\t'+str(item['page']));runfont(rr,'Song',11)
                # Add internal hyperlink by wrapping existing content runs.
                hyp=e('w:hyperlink',**{'w:anchor':item['anchor']})
                for rr in list(pp._p):
                    if rr.tag==qn('w:r'):pp._p.remove(rr);hyp.append(rr)
                pp._p.append(hyp)
            else:paragraph(item,'Title' if kind in ['cover_title','toc_title'] else ('Subtitle' if kind=='cover_sub' else None),gap)
            prev_bottom=item['top']+len(item['lines'])*item['leading']
            continue
        if kind=='table':
            tb=doc.add_table(rows=0,cols=len(item['widths']));tb.autofit=False;tb.alignment=WD_TABLE_ALIGNMENT.CENTER
            for co,w in zip(tb.columns,item['widths']):co.width=Pt(w)
            pr=tb._tbl.tblPr;bor=e('w:tblBorders')
            for side in ['top','left','bottom','right','insideH','insideV']:bor.append(e('w:'+side,**{'w:val':'single','w:sz':'4','w:color':'D9D9D9'}))
            pr.append(bor);mar=e('w:tblCellMar')
            for side in ['top','left','bottom','right']:mar.append(e('w:'+side,**{'w:w':int(item['pad']*20),'w:type':'dxa'}))
            pr.append(mar)
            for ri,row in enumerate(item['rows']):
                rw=tb.add_row();tr=rw._tr.get_or_add_trPr();tr.append(e('w:cantSplit'));tr.append(e('w:trHeight',**{'w:val':int(row['h']*20),'w:hRule':'atLeast'}))
                if row['ishead']:tr.append(e('w:tblHeader'))
                for cell,w,ls in zip(rw.cells,item['widths'],row['cells']):
                    cell.width=Pt(w);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;cell._tc.get_or_add_tcPr().append(e('w:shd',**{'w:fill':'E7EEF4' if row['ishead'] else ('F7F9FB' if ri%2==0 else 'FFFFFF')}))
                    pp=cell.paragraphs[0];fmt=pp.paragraph_format;fmt.line_spacing=Pt(item['leading']);fmt.space_before=Pt(0);fmt.space_after=Pt(0);fmt.widow_control=False;fmt.keep_with_next=False
                    pp._p.get_or_add_pPr().append(e('w:snapToGrid',**{'w:val':'0'}))
                    for j,line in enumerate(ls):
                        if j:pp.add_run().add_break()
                        rr=pp.add_run(line);runfont(rr,'Hei' if row['ishead'] else 'Song',item['size'])
            pp=doc.add_paragraph();pp.paragraph_format.line_spacing=Pt(6);pp.paragraph_format.space_after=Pt(0);rr=pp.add_run('');runfont(rr,'Song',1)
        elif kind=='figure':
            pp=doc.add_paragraph();pp.paragraph_format.line_spacing=1;pp.paragraph_format.space_after=Pt(0);pp.paragraph_format.keep_with_next=True
            rr=pp.add_run();rr.add_picture(item['path'],width=Pt(CW));rr._r.xpath('.//wp:docPr')[0].set('descr',item['caption'])
            pp=doc.add_paragraph(item['caption']);pp.alignment=WD_ALIGN_PARAGRAPH.CENTER;pp.paragraph_format.line_spacing=Pt(19);pp.paragraph_format.space_after=Pt(6)
            for rr in pp.runs:runfont(rr,'Song',10)
        elif kind=='link':
            pp=doc.add_paragraph();pp.paragraph_format.line_spacing=Pt(14);pp.paragraph_format.space_after=Pt(1)
            from docx.opc.constants import RELATIONSHIP_TYPE as RT
            rid=pp.part.relate_to(item['url'],RT.HYPERLINK,is_external=True);hyp=e('w:hyperlink',**{'r:id':rid});rr=e('w:r');rp=e('w:rPr');rp.append(e('w:color',**{'w:val':'1F4E79'}));rp.append(e('w:sz',**{'w:val':'20'}));rp.append(e('w:rFonts',**{'w:ascii':'宋体','w:eastAsia':'宋体'}));rr.append(rp);tt=e('w:t');tt.text='原始资料';rr.append(tt);hyp.append(rr);pp._p.append(hyp)
        else:paragraph(item,{'h1':'Heading 1','h2':'Heading 2','h3':'Heading 3'}.get(kind))
for pp in doc.paragraphs:
    pr=pp._p.get_or_add_pPr();pr.append(e('w:autoSpaceDE',**{'w:val':'0'}));pr.append(e('w:autoSpaceDN',**{'w:val':'0'}))
for tb in doc.tables:
    for row in tb.rows:
        for cell in row.cells:
            for pp in cell.paragraphs:
                pr=pp._p.get_or_add_pPr();pr.append(e('w:autoSpaceDE',**{'w:val':'0'}));pr.append(e('w:autoSpaceDN',**{'w:val':'0'}))
doc.core_properties.title=D['title'];doc.core_properties.author='';doc.core_properties.subject='智慧农业本体论数据库建设与实施'
doc.save(OUT)
(BASE/'pagination.json').write_text(json.dumps({'pages':len(PAGES),'toc':toc,'layout':PAGES,'render_method':'Explicit line and page layout; native Word rendering unavailable'},ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'docx':str(OUT),'preview':str(PDF),'pages':len(PAGES),'chapter_pages':toc},ensure_ascii=False))
