import json, re, math
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

BASE=Path(__file__).resolve().parent
ROOT=BASE.parent
DATA=json.loads((BASE/'plan-content.json').read_text(encoding='utf-8'))
OUT=ROOT/'交付文档'/'梁子湖智慧农业本体论数据库总体建设方案.docx'
FONT='C:/Windows/Fonts/msyh.ttc'
FONTSIM='C:/Windows/Fonts/simsun.ttc'

def diagram(name):
    if name=='architecture':
        im=Image.new('RGB',(1600,1080),'white'); dr=ImageDraw.Draw(im)
        boxes=[(70,40,1530,160,'多源数据接入','视频与图像  天气  土壤水质  遥感光谱  报告  手工与 API'),(70,220,1530,340,'原始资产与接收记录','内容哈希  来源版本  接收时间  私有访问  精确定位'),(70,400,1530,540,'处理与语义治理','解析与特征  身份消歧  时空对齐  单位映射  约束与冲突'),(70,600,1530,740,'权威事实与证据存储','现有业务表  观测与事件  断言与限定关系  处理运行'),(70,800,770,990,'图与检索视图','邻域  证据链  时空回放'),(830,800,1530,990,'推理信息包','前提  反证  缺失  版本')]
        for x1,y1,x2,y2,a,b in boxes:
            dr.rounded_rectangle((x1,y1,x2,y2),radius=12,fill='#F2F6F9',outline='#758694',width=2)
            for text,y,size in [(a,y1+25,38),(b,y1+80,29)]:
                f=ImageFont.truetype(FONT,size); bb=dr.textbbox((0,0),text,font=f)
                dr.text(((x1+x2-(bb[2]-bb[0]))/2,y),text,font=f,fill='#17212B')
        for y1,y2 in [(165,213),(345,393),(545,593)]:
            dr.line((800,y1,800,y2),fill='#455A64',width=4);dr.polygon([(790,y2-12),(810,y2-12),(800,y2)],fill='#455A64')
        for x in (420,1180):
            dr.line((x,745,x,792),fill='#455A64',width=4);dr.polygon([(x-10,780),(x+10,780),(x,792)],fill='#455A64')
        dr.text((70,1020),'横贯各层  概念与身份  关系语义  时空与事件  观测与证据  约束与推理',font=ImageFont.truetype(FONT,28),fill='#283640')
    else:
        im=Image.new('RGB',(1600,880),'white');dr=ImageDraw.Draw(im)
        nodes=[(40,60,480,230,'原始证据','报文  图像区域  报告段落'),(580,60,1020,230,'标准记录','观测  事件  来源陈述'),(1120,60,1560,230,'待解释现象','真实变化或测量问题'),(40,480,480,670,'处理与映射运行','方法  参数  版本  身份映射'),(580,480,1020,670,'推理运行','实际前提  规则  快照'),(1120,480,1560,670,'候选成因','机制  支持  反证  缺失')]
        for x1,y1,x2,y2,a,b in nodes:
            dr.rounded_rectangle((x1,y1,x2,y2),radius=10,fill='#F2F6F9',outline='#758694',width=2)
            for text,y,size in [(a,y1+32,37),(b,y1+101,25)]:
                f=ImageFont.truetype(FONT,size);bb=dr.textbbox((0,0),text,font=f);dr.text(((x1+x2-bb[2])/2,y),text,font=f,fill='#17212B')
        def arrow(x1,y1,x2,y2,label,lx,ly):
            dr.line((x1,y1,x2,y2),fill='#455A64',width=4)
            a=math.atan2(y2-y1,x2-x1);s=13
            dr.polygon([(x2,y2),(x2-s*math.cos(a-.55),y2-s*math.sin(a-.55)),(x2-s*math.cos(a+.55),y2-s*math.sin(a+.55))],fill='#455A64')
            dr.text((lx,ly),label,font=ImageFont.truetype(FONT,27),fill='#283640')
        arrow(480,145,570,145,'提取',488,90);arrow(1020,145,1110,145,'描述',1028,90)
        arrow(260,475,260,240,'记录生成依据',95,337)
        arrow(800,240,800,470,'采用的前提',820,337)
        arrow(1340,240,1340,470,'提出并检验',1360,337)
        arrow(1030,575,1110,575,'产生',1030,520)
        dr.text((55,750),'来源链回答从哪里来  依据链回答如何得出  成因链检验现实中的可能原因',font=ImageFont.truetype(FONT,30),fill='#283640')
        dr.text((55,815),'支持或反驳关系也记录评估方法与版本  新证据可触发重新评估',font=ImageFont.truetype(FONT,28),fill='#283640')
    path=BASE/(name+'.png');im.save(path);return path

def el(tag,**attrs):
    e=OxmlElement(tag)
    for k,v in attrs.items():e.set(qn(k),str(v))
    return e

def setfont(style,name,size,east='宋体',bold=False):
    style.font.name=name;style.font.size=Pt(size);style.font.bold=bold;style.font.color.rgb=RGBColor(0,0,0)
    rpr=style.element.get_or_add_rPr();fonts=rpr.find(qn('w:rFonts'))
    if fonts is None:fonts=el('w:rFonts');rpr.append(fonts)
    for attr,val in [('w:ascii',name),('w:hAnsi',name),('w:eastAsia',east)]:fonts.set(qn(attr),val)
    lang=el('w:lang',**{'w:val':'en-US','w:eastAsia':'zh-CN'});rpr.append(lang)

def no_border(p):
    pr=p._p.get_or_add_pPr();bd=pr.find(qn('w:pBdr'))
    if bd is not None:pr.remove(bd)

def field(p,instr):
    r=p.add_run();r._r.append(el('w:fldChar',**{'w:fldCharType':'begin'}))
    x=el('w:instrText');x.set(qn('xml:space'),'preserve');x.text=' '+instr+' ';r._r.append(x)
    r._r.append(el('w:fldChar',**{'w:fldCharType':'separate'}));r._r.append(el('w:t'));r._r.append(el('w:fldChar',**{'w:fldCharType':'end'}))

def hyperlink(p,label,url):
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    rid=p.part.relate_to(url,RT.HYPERLINK,is_external=True)
    hyp=el('w:hyperlink',**{'r:id':rid});r=el('w:r');pr=el('w:rPr');pr.append(el('w:color',**{'w:val':'1F4E79'}));pr.append(el('w:u',**{'w:val':'single'}));r.append(pr);txt=el('w:t');txt.text=label;r.append(txt);hyp.append(r);p._p.append(hyp)

doc=Document();sec=doc.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=Cm(2.0);sec.bottom_margin=Cm(2.0);sec.left_margin=Cm(2.1);sec.right_margin=Cm(2.1)
sec.header_distance=Cm(.9);sec.footer_distance=Cm(.9);sec.different_first_page_header_footer=True
styles=doc.styles
setfont(styles['Normal'],'Calibri',11,'宋体');styles['Normal'].paragraph_format.line_spacing=1.3;styles['Normal'].paragraph_format.space_after=Pt(6)
setfont(styles['Title'],'Calibri',25,'微软雅黑',True);styles['Title'].paragraph_format.space_after=Pt(18)
setfont(styles['Subtitle'],'Calibri',14,'微软雅黑');styles['Subtitle'].paragraph_format.space_after=Pt(20)
for name,size in [('Heading 1',17),('Heading 2',13),('Heading 3',11.5)]:
    setfont(styles[name],'Calibri',size,'微软雅黑',True);styles[name].paragraph_format.space_before=Pt(14);styles[name].paragraph_format.space_after=Pt(8);styles[name].paragraph_format.keep_with_next=True
styles['Heading 1'].paragraph_format.page_break_before=True
for name in ['Caption','List Bullet']:
    setfont(styles[name],'Calibri',10 if name=='Caption' else 11,'宋体')
styles['Caption'].paragraph_format.space_after=Pt(9)
for name in ['Code','Reference','TOC 1']:
    if name not in styles:styles.add_style(name,1)
setfont(styles['Code'],'Consolas',9,'微软雅黑');styles['Code'].paragraph_format.line_spacing=1.05;styles['Code'].paragraph_format.space_after=Pt(0);styles['Code'].paragraph_format.space_before=Pt(0)
setfont(styles['Reference'],'Calibri',10,'宋体');styles['Reference'].paragraph_format.line_spacing=1.2;styles['Reference'].paragraph_format.space_after=Pt(8)
setfont(styles['TOC 1'],'Calibri',11,'宋体');styles['TOC 1'].paragraph_format.space_after=Pt(9)
for s in styles:
    if s.type==1 and s.name.startswith('Heading'):
        s.font.color.rgb=RGBColor(0,0,0)
# Page numbers are useful in a long implementation plan.
f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=f.add_run('梁子湖智慧农业本体论数据库  ·  ');r.font.size=Pt(9)
field(f,'PAGE');r=f.add_run(' / ');r.font.size=Pt(9);field(f,'NUMPAGES')
# Cover and executive orientation.
p0=doc.add_paragraph(DATA['title'],style='Title');p0.paragraph_format.space_before=Pt(64);no_border(p0)
doc.add_paragraph(DATA['subtitle'],style='Subtitle')
p0=doc.add_paragraph('版本 '+DATA['version']+'     '+DATA['date']);p0.paragraph_format.space_after=Pt(24)
doc.add_paragraph('建设重点',style='Heading 2')
doc.add_paragraph('构建从信源接收、语义对齐、事实与证据管理，到代码层推理输入的完整数据基础。以虾稻共作为试点，统一概念与身份、关系语义、时空与事件、观测与证据、约束与推理五个方面，并将可溯源性贯穿每一步。')
doc.add_paragraph('建议保留 PostgreSQL 作为权威存储，在现有梁子湖农场平台上新增语义、证据及推理接口模块。属性图与 RDF 作为可重建和可交换的表示。首期完成数据与证据闭环，生产自动决策及设备执行按独立边界扩展。')
doc.add_paragraph('阅读使用',style='Heading 2')
doc.add_paragraph('第一至三章说明目标、架构和平台集成；第四至六章定义语义、知识图及标准交换；第七至十一章说明数据处理、冲突、溯源、存储与推理接口；第十二至十三章给出实施和验收方法。附录提供契约、SQL、JSON、RDF 示例及交付清单。')
doc.add_page_break();pp=doc.add_paragraph('目录',style='Title');pp.paragraph_format.space_before=Pt(0)
pp=doc.add_paragraph();field(pp,'TOC \\o "1-1" \\h \\z \\u')
# Table builder with explicit widths and borders.
def table(block):
    tb=doc.add_table(rows=1,cols=len(block['headers']));tb.alignment=WD_TABLE_ALIGNMENT.CENTER;tb.autofit=False
    widths=block.get('widths') or [100/len(block['headers'])]*len(block['headers'])
    usable=16.8
    for col,w in zip(tb.columns,widths):col.width=Cm(usable*w/100)
    props=tb._tbl.tblPr;borders=el('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:borders.append(el('w:'+edge,**{'w:val':'single','w:sz':'4','w:color':'D9D9D9'}))
    props.append(borders)
    mar=el('w:tblCellMar')
    for edge,value in [('top',95),('bottom',95),('left',105),('right',105)]:mar.append(el('w:'+edge,**{'w:w':value,'w:type':'dxa'}))
    props.append(mar)
    allrows=[block['headers']]+block['rows']
    for idx,rowdata in enumerate(allrows):
        row=tb.rows[0] if idx==0 else tb.add_row();rp=row._tr.get_or_add_trPr();rp.append(el('w:cantSplit'))
        if idx==0:rp.append(el('w:tblHeader'))
        for j,(cell,value) in enumerate(zip(row.cells,rowdata)):
            cell.width=Cm(usable*widths[j]/100);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text=str(value);cp=cell._tc.get_or_add_tcPr();cp.append(el('w:shd',**{'w:fill':'E7EEF4' if idx==0 else ('F7F9FB' if idx%2==0 else 'FFFFFF')}))
            pp=cell.paragraphs[0];pp.paragraph_format.space_after=Pt(0);pp.paragraph_format.line_spacing=1.16;pp.paragraph_format.keep_with_next=(idx==0)
            pp.alignment=WD_ALIGN_PARAGRAPH.LEFT
            for run in pp.runs:
                run.font.name='Calibri';run.font.size=Pt(9.5);run.font.bold=(idx==0);run.font.color.rgb=RGBColor(0,0,0)
                run._element.get_or_add_rPr().append(el('w:rFonts',**{'w:ascii':'Calibri','w:hAnsi':'Calibri','w:eastAsia':'宋体'}))
    pp=doc.add_paragraph();pp.paragraph_format.space_after=Pt(2);pp.paragraph_format.space_before=Pt(0);pp.paragraph_format.line_spacing=Pt(3);pp.add_run().font.size=Pt(3)

for chapter in DATA['chapters']:
    doc.add_paragraph(chapter['title'],style='Heading 1')
    for x in chapter['blocks']:
        kind=x['type']
        if kind=='p':doc.add_paragraph(x['text'])
        elif kind=='h2':doc.add_paragraph(x['text'],style='Heading 2')
        elif kind=='h3':doc.add_paragraph(x['text'],style='Heading 3')
        elif kind=='table':table(x)
        elif kind=='bullets':
            for item in x['items']:
                pp=doc.add_paragraph(item,style='List Bullet');pp.paragraph_format.space_after=Pt(5)
        elif kind=='code':
            for line in x['text'].splitlines():doc.add_paragraph(line,style='Code')
            doc.add_paragraph().paragraph_format.space_after=Pt(3)
        elif kind=='figure':
            pp=doc.add_paragraph();pp.alignment=WD_ALIGN_PARAGRAPH.CENTER;pp.paragraph_format.keep_with_next=True
            run=pp.add_run();run.add_picture(str(diagram(x['name'])),width=Cm(16.5))
            run._r.xpath('.//wp:docPr')[0].set('descr',x['caption'])
            pp=doc.add_paragraph(x['caption'],style='Caption');pp.alignment=WD_ALIGN_PARAGRAPH.CENTER
        elif kind=='reference':
            pp=doc.add_paragraph(style='Reference');rr=pp.add_run('['+x['id']+'] '+x['title']+'。');rr.bold=True
            pp.add_run(x['note'])
            if x['url']:pp.add_run(' ');hyperlink(pp,'原始资料',x['url'])
settings=doc.settings.element;settings.append(el('w:updateFields',**{'w:val':'true'}))
doc.core_properties.title=DATA['title'];doc.core_properties.subject='农业本体 证据治理 多模态数据 推理接口'
doc.core_properties.author='';doc.core_properties.keywords='梁子湖 智慧农业 本体论 虾稻共作 可溯源'
doc.save(OUT)
# Content checks are independent of visual QA.
text=json.dumps(DATA,ensure_ascii=False)
checks={
 'chapters':len(DATA['chapters']),
 'tables':sum(b['type']=='table' for c in DATA['chapters'] for b in c['blocks']),
 'figures':sum(b['type']=='figure' for c in DATA['chapters'] for b in c['blocks']),
 'references':len(DATA['references']),
 'json_examples':0,
 'required_topics_present':all(s in text for s in ['概念边界','身份规则','关系语义','事件','断言','SHACL','RDF','冲突','反证','outbox','回滚','校准','SourceContract','EvidenceBundle']),
 'bad_unicode': bool(re.search('[\ufffd\uff80-\uffef]',text))
}
for c in DATA['chapters']:
 for x in c['blocks']:
  if x['type']=='code' and x['text'].lstrip().startswith('{'):
   json.loads(x['text']);checks['json_examples']+=1
assert checks['required_topics_present'] and not checks['bad_unicode']
(BASE/'content-checks.json').write_text(json.dumps(checks,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'file':str(OUT),**checks},ensure_ascii=False))
