"""Pure, bounded file decoding under immutable server-resolved upload identity."""
import os
import sys

# Bootstrap precedes libraries, input and decoding. Parent attaches Windows
# hard memory confinement before sending input. This is not an OS sandbox.
if __name__ == "__main__":
    # openpyxl optionally imports installed NumPy for scalar type recognition.
    # No numerical worker pool is needed for file decoding inside this cap.
    os.environ["OPENBLAS_NUM_THREADS"] = "1"
    os.environ["OMP_NUM_THREADS"] = "1"
    if sys.platform.startswith("linux"):
        try:
            import resource
            memory = int(os.environ["LZO_QUERY_CHILD_MEMORY_BYTES"])
            resource.setrlimit(resource.RLIMIT_AS, (memory, memory))
        except (ImportError, KeyError, OSError, ValueError):
            sys.stdout.write('{"error":"FILE_CAPABILITY_UNAVAILABLE"}')
            sys.exit(0)
    def confinement(event, args):
        if event.startswith("socket.") or event in {"urllib.Request", "subprocess.Popen", "os.system", "os.exec", "os.posix_spawn"}:
            raise PermissionError("File parser external access denied")
    sys.addaudithook(confinement)

import base64
import codecs
import csv
import hashlib
import io
import json
import math
import re
import zipfile
from dataclasses import asdict, dataclass
from datetime import date, datetime, time, timezone
from uuid import uuid5
from zoneinfo import ZoneInfo
from lzo.ingestion.schemas import ParsedFile, ParseIssue, SourceChange, SourcePosition, UploadBinding
from lzo.ingestion.datetimes import source_datetime


@dataclass(frozen=True)
class ParserLimits:
    max_bytes: int = 8*1024*1024
    max_rows: int = 10000
    max_columns: int = 256
    max_cells: int = 100000
    max_text: int = 65536
    max_depth: int = 32
    max_expanded_bytes: int = 32*1024*1024
    max_members: int = 256
    max_ratio: int = 200
    max_output_bytes: int = 16*1024*1024
    timeout_ms: int = 20000
    memory_mb: int = 512

    def __post_init__(self):
        if any(v <= 0 for v in asdict(self).values()):
            raise ValueError("File budgets must be positive")

    @classmethod
    def from_settings(cls, settings):
        return cls(**{name:getattr(settings, "ingestion_"+name) for name in cls.__dataclass_fields__})


class InvalidFile(Exception):
    def __init__(self, code, locator="$", field=None):
        self.code, self.locator, self.field = code, locator, field


def issue(code, locator="$", field=None):
    return ParseIssue(code=code, locator=locator, field_path=field, message="File does not satisfy the controlled parsing contract")


def incomplete(code, locator="$", field=None):
    return ParsedFile(changes=(), issues=(issue(code, locator, field),), complete=False)


def canonical(value):
    return json.dumps(value, sort_keys=True, separators=(",", ":"), ensure_ascii=False, allow_nan=False)


def path_value(value, path):
    if not path:
        return value
    parts = path[1:].split("/") if path.startswith("/") else [path]
    try:
        for part in parts:
            part = part.replace("~1", "/").replace("~0", "~")
            value = value[int(part)] if isinstance(value, list) else value[part]
        return value
    except (KeyError, TypeError, IndexError, ValueError):
        raise InvalidFile("MISSING_PATH", field=path) from None


def normalize(value, options, limits, locator, depth=0, work=None):
    work = [0] if work is None else work
    work[0] += 1
    if work[0] > limits.max_cells:
        raise InvalidFile("CELL_LIMIT", locator)
    if depth > limits.max_depth:
        raise InvalidFile("DEPTH_LIMIT", locator)
    if isinstance(value, str):
        if len(value.encode("utf-8")) > limits.max_text:
            raise InvalidFile("TEXT_LIMIT", locator)
        return value
    if isinstance(value, datetime):
        if not options.source_timezone:
            raise InvalidFile("TIMEZONE_REQUIRED", locator)
        try:
            return source_datetime(value,options.source_timezone).isoformat()
        except (ValueError, KeyError):
            raise InvalidFile("INVALID_TIMEZONE", locator) from None
    if isinstance(value, (date, time)):
        return value.isoformat()
    if value is None or isinstance(value, (bool, int)):
        return value
    if isinstance(value, float):
        if not math.isfinite(value):
            raise InvalidFile("NONFINITE_NUMBER", locator)
        return value
    if isinstance(value, dict):
        if len(value) > limits.max_columns:
            raise InvalidFile("COLUMN_LIMIT", locator)
        return {normalize(k,options,limits,locator,depth+1,work): normalize(v,options,limits,locator,depth+1,work) for k,v in value.items()}
    if isinstance(value, (list, tuple)):
        if len(value) > limits.max_cells:
            raise InvalidFile("CELL_LIMIT", locator)
        return [normalize(v,options,limits,locator,depth+1,work) for v in value]
    raise InvalidFile("INVALID_CELL_TYPE", locator)


def headers(row, limits, locator):
    if not row or len(row) > limits.max_columns:
        raise InvalidFile("COLUMN_LIMIT", locator)
    if any(not isinstance(v,str) or not v for v in row):
        raise InvalidFile("INVALID_HEADER", locator)
    if len(set(row)) != len(row):
        raise InvalidFile("DUPLICATE_HEADER", locator)
    return row


def csv_rows(data, options, limits):
    try:
        codec = codecs.lookup(options.encoding)
        if codec.name not in {"utf-8", "utf-8-sig", "utf-16", "utf-16-le", "utf-16-be", "gb18030", "gbk", "iso8859-1", "cp1252"}:
            raise LookupError()
        text = data.decode(options.encoding, errors="strict")
    except (UnicodeError, LookupError):
        raise InvalidFile("INVALID_ENCODING") from None
    csv.field_size_limit(limits.max_text)
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=options.delimiter, strict=True)
    names = None
    try:
        for row_number, row in enumerate(reader, 1):
            if row_number < options.header_row:
                continue
            locator = f"row:{row_number}"
            if names is None:
                names = headers(row, limits, locator)
                continue
            if len(row) != len(names):
                raise InvalidFile("ROW_WIDTH", locator)
            yield dict(zip(names,row)), None, locator
        if names is None:
            raise InvalidFile("HEADER_MISSING")
    except csv.Error:
        raise InvalidFile("INVALID_CSV", f"row:{reader.line_num}") from None


def checked_zip(data, limits):
    from defusedxml import ElementTree
    try:
        archive = zipfile.ZipFile(io.BytesIO(data))
        infos = archive.infolist()
        if len(infos) > limits.max_members:
            raise InvalidFile("ZIP_MEMBER_LIMIT")
        names, total = set(), 0
        if not {"[Content_Types].xml", "xl/workbook.xml"} <= {i.filename for i in infos}:
            raise InvalidFile("INVALID_XLSX")
        for entry in infos:
            name = entry.filename
            locator = "package:"+str(len(names)+1)
            if name in names or name.startswith(("/","\\")) or ".." in name.split("/") or "\\" in name or entry.flag_bits & 1:
                raise InvalidFile("INVALID_ZIP_MEMBER", locator)
            names.add(name)
            if not name.endswith((".xml", ".rels")) or entry.compress_type not in (zipfile.ZIP_STORED, zipfile.ZIP_DEFLATED):
                raise InvalidFile("UNSUPPORTED_XLSX_PART", locator)
            if entry.file_size > limits.max_expanded_bytes or entry.file_size > max(1,entry.compress_size)*limits.max_ratio:
                raise InvalidFile("ZIP_EXPANSION_LIMIT", locator)
            expanded = bytearray()
            with archive.open(entry) as stream:
                while chunk := stream.read(65536):
                    total += len(chunk)
                    expanded.extend(chunk)
                    if total > limits.max_expanded_bytes or len(expanded) > max(1,entry.compress_size)*limits.max_ratio:
                        raise InvalidFile("ZIP_EXPANSION_LIMIT", locator)
            root = ElementTree.fromstring(expanded)
            work, nodes = [(root,0)], 0
            while work:
                node, depth = work.pop(); nodes += 1
                if depth > limits.max_depth or nodes > limits.max_cells*8:
                    raise InvalidFile("XML_LIMIT", locator)
                if node.attrib.get("TargetMode") == "External" or any("macroEnabled" in v or "vbaProject" in v for v in node.attrib.values()):
                    raise InvalidFile("EXTERNAL_OR_MACRO_REFERENCE", locator)
                if node.text and len(node.text.encode()) > limits.max_text:
                    raise InvalidFile("TEXT_LIMIT", locator)
                work.extend((child,depth+1) for child in node)
        return archive
    except InvalidFile:
        raise
    except Exception:
        raise InvalidFile("INVALID_XLSX") from None


def checked_worksheet(data, worksheet_path, options, limits):
    """Check stored coordinates before read-only reconstruction can lose cells.

    The package has already passed actual expansion/XML budgets. The pinned
    openpyxl reader resolves the selected sheet's path from workbook relations;
    use that same path rather than guessing a sheet number from its title.
    """
    from defusedxml import ElementTree
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        root=ElementTree.fromstring(archive.read(worksheet_path))
    ns="{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    locator="sheet:"+options.sheet
    sections=root.findall(ns+"sheetData")
    if root.tag!=ns+"worksheet" or len(sections)>1:
        raise InvalidFile("INVALID_XLSX",locator)
    rows=sections[0].findall(ns+"row") if sections else []
    cells=[cell for row in rows for cell in row.findall(ns+"c")]
    # Reject misplaced nodes that a streaming reader might still interpret.
    if list(root.iter(ns+"row"))!=rows or list(root.iter(ns+"c"))!=cells:
        raise InvalidFile("INVALID_COORDINATE",locator)
    if len(cells)>limits.max_cells:
        raise InvalidFile("CELL_LIMIT",locator)
    previous_row=0
    for row in rows:
        raw_row=row.get("r","")
        if not re.fullmatch(r"[1-9][0-9]{0,6}",raw_row):
            raise InvalidFile("INVALID_COORDINATE",locator+"!row:unknown")
        row_number=int(raw_row)
        row_locator=locator+"!row:"+str(row_number)
        if row_number>min(1048576,limits.max_rows+options.header_row):
            raise InvalidFile("DIMENSION_LIMIT",row_locator)
        if row_number<=previous_row:
            raise InvalidFile("INVALID_COORDINATE",row_locator)
        previous_row=row_number
        previous_column=0
        for cell in row.findall(ns+"c"):
            coordinate=cell.get("r","")
            match=re.fullmatch(r"([A-Z]{1,3})([1-9][0-9]{0,6})",coordinate)
            if match is None:
                raise InvalidFile("INVALID_COORDINATE",row_locator)
            cell_locator=locator+"!"+coordinate
            column=0
            for char in match[1]: column=column*26+ord(char)-ord("A")+1
            if column>min(16384,limits.max_columns):
                raise InvalidFile("DIMENSION_LIMIT",cell_locator)
            if int(match[2])!=row_number or column<=previous_column:
                raise InvalidFile("INVALID_COORDINATE",cell_locator)
            previous_column=column
            if next(cell.iter(ns+"f"),None) is not None:
                raise InvalidFile("FORMULA_CELL",cell_locator)
            if cell.get("t")=="e":
                raise InvalidFile("ERROR_CELL",cell_locator)


def xlsx_rows(data, options, limits):
    from openpyxl import load_workbook
    checked_zip(data, limits).close()
    if not options.sheet:
        raise InvalidFile("SHEET_REQUIRED")
    try:
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
        try:
            if len(book.sheetnames) > limits.max_members or options.sheet not in book.sheetnames:
                raise InvalidFile("INVALID_SHEET")
            sheet = book[options.sheet]
            if (sheet.max_row or 0) > limits.max_rows+options.header_row or (sheet.max_column or 0) > limits.max_columns:
                raise InvalidFile("DIMENSION_LIMIT", "sheet:"+options.sheet)
            checked_worksheet(data,sheet._worksheet_path,options,limits)
            names, count = None, 0
            # Reset dimensions after validating declarations: maliciously small
            # cached dimensions must not hide actual cells or formulas.
            sheet.reset_dimensions()
            for row_number, row in enumerate(sheet.iter_rows(),1):
                if row_number > limits.max_rows+options.header_row or len(row)>limits.max_columns:
                    raise InvalidFile("DIMENSION_LIMIT", "sheet:"+options.sheet)
                count += len(row)
                if count > limits.max_cells:
                    raise InvalidFile("CELL_LIMIT", "sheet:"+options.sheet)
                for cell in row:
                    if cell.data_type == "f":
                        raise InvalidFile("FORMULA_CELL", f"sheet:{options.sheet}!{cell.coordinate}")
                    if cell.data_type == "e":
                        raise InvalidFile("ERROR_CELL", f"sheet:{options.sheet}!{cell.coordinate}")
                if row_number < options.header_row:
                    continue
                locator=f"sheet:{options.sheet}!row:{row_number}"
                values=[cell.value for cell in row]
                if names is None:
                    names=headers(values,limits,locator)
                    continue
                if len(values)>len(names):
                    raise InvalidFile("ROW_WIDTH",locator)
                values.extend([None]*(len(names)-len(values)))
                yield dict(zip(names,values)), None, locator
            if names is None:
                raise InvalidFile("HEADER_MISSING")
        finally:
            book.close()
    except InvalidFile:
        raise
    except Exception:
        raise InvalidFile("INVALID_XLSX") from None


def json_rows(data, options, limits):
    def pairs(items):
        result={}
        for key,value in items:
            if key in result:
                raise InvalidFile("DUPLICATE_JSON_KEY")
            result[key]=value
        return result
    def nonfinite(value):
        raise InvalidFile("NONFINITE_NUMBER")
    try:
        if options.encoding.lower().replace("_","-") not in {"utf-8","utf-8-sig"}:
            raise InvalidFile("INVALID_ENCODING")
        root=json.loads(data.decode(options.encoding,"strict"), object_pairs_hook=pairs, parse_constant=nonfinite)
    except UnicodeError:
        raise InvalidFile("INVALID_ENCODING") from None
    except (ValueError, RecursionError):
        raise InvalidFile("INVALID_JSON") from None
    root=normalize(root, options, limits, "$")
    crs=None
    if options.format=="geojson":
        if not isinstance(root,dict) or root.get("type")!="FeatureCollection" or options.records_path:
            raise InvalidFile("INVALID_GEOJSON")
        crs=root.get("crs")
        records=root.get("features")
    else:
        records=path_value(root,options.records_path)
    if not isinstance(records,list):
        raise InvalidFile("RECORDS_NOT_ARRAY")
    for ordinal, record in enumerate(records):
        locator=f"{options.records_path or '$'}/{ordinal}"
        geometry=None
        if options.format=="geojson":
            if not isinstance(record,dict) or record.get("type")!="Feature" or not isinstance(record.get("properties"),dict):
                raise InvalidFile("INVALID_GEOJSON", locator)
            geometry=record.get("geometry")
            if geometry is not None:
                if not isinstance(geometry,dict) or geometry.get("type") not in {"Point","MultiPoint","LineString","MultiLineString","Polygon","MultiPolygon","GeometryCollection"}:
                    raise InvalidFile("INVALID_GEOMETRY",locator)
                def geometry_check(g):
                    if not isinstance(g,dict): raise InvalidFile("INVALID_GEOMETRY",locator)
                    if g.get("type")=="GeometryCollection":
                        if not isinstance(g.get("geometries"),list): raise InvalidFile("INVALID_GEOMETRY",locator)
                        for child in g["geometries"]: geometry_check(child)
                        return
                    dimensions={"Point":0,"MultiPoint":1,"LineString":1,"MultiLineString":2,"Polygon":2,"MultiPolygon":3}
                    if g.get("type") not in dimensions: raise InvalidFile("INVALID_GEOMETRY",locator)
                    def coordinates(v,depth):
                        if not isinstance(v,list) or not v: raise InvalidFile("INVALID_GEOMETRY",locator)
                        if depth:
                            for part in v: coordinates(part,depth-1)
                        elif len(v)<2 or any(isinstance(n,bool) or not isinstance(n,(int,float)) or not math.isfinite(n) for n in v):
                            raise InvalidFile("INVALID_GEOMETRY",locator)
                    coordinates(g.get("coordinates"),dimensions[g["type"]])
                    if g["type"] in {"LineString","MultiLineString"}:
                        lines=[g["coordinates"]] if g["type"]=="LineString" else g["coordinates"]
                        for line in lines:
                            if len(line)<2: raise InvalidFile("INVALID_GEOMETRY",locator)
                    if g["type"] in {"Polygon","MultiPolygon"}:
                        polygons=[g["coordinates"]] if g["type"]=="Polygon" else g["coordinates"]
                        for polygon in polygons:
                            for ring in polygon:
                                if len(ring)<4 or ring[0]!=ring[-1]: raise InvalidFile("INVALID_GEOMETRY",locator)
                geometry_check(geometry)
                geometry=dict(geometry)
                if crs is not None: geometry["crs"]=crs
                if options.source_crs is not None: geometry["source_crs"]=options.source_crs
            record=record["properties"]
        if not isinstance(record,dict):
            raise InvalidFile("RECORD_NOT_OBJECT",locator)
        yield record, geometry, locator


def decode(data, binding, limits):
    changes=[]
    try:
        options=binding.options
        if options.header_row > limits.max_rows or (options.format in {"csv","xlsx"} and options.records_path):
            raise InvalidFile("INVALID_PARSE_CONFIGURATION")
        reader = csv_rows if options.format=="csv" else xlsx_rows if options.format=="xlsx" else json_rows
        seen, cells, output_size = set(), 0, 0
        for ordinal,(raw,geometry,locator) in enumerate(reader(data,options,limits),1):
            if ordinal>limits.max_rows:
                raise InvalidFile("ROW_LIMIT",locator)
            raw=normalize(raw,options,limits,locator)
            cells+=len(raw)
            if cells>limits.max_cells: raise InvalidFile("CELL_LIMIT",locator)
            keys=[path_value(raw,path) for path in options.key_paths]
            if any(k is None or isinstance(k,(list,dict,bool)) or k=="" for k in keys):
                raise InvalidFile("INVALID_KEY",locator)
            key=canonical(keys)
            if key in seen: raise InvalidFile("DUPLICATE_KEY",locator)
            seen.add(key)
            temporal={"valid_from":None,"valid_to":None,"observed_at":None}
            if options.temporal_fields is not None:
                for name in temporal:
                    path=getattr(options.temporal_fields,name+"_path")
                    if path is None: continue
                    try:
                        value=path_value(raw,path)
                        temporal[name]=None if value is None else source_datetime(value,options.source_timezone)
                    except InvalidFile as error:
                        raise InvalidFile(error.code,locator,path) from None
                    except (ValueError,TypeError,OverflowError):
                        raise InvalidFile("INVALID_TEMPORAL_VALUE",locator,path) from None
                if temporal["valid_from"] is not None and temporal["valid_to"] is not None and temporal["valid_to"]<=temporal["valid_from"]:
                    raise InvalidFile("INVALID_TEMPORAL_INTERVAL",locator,options.temporal_fields.valid_to_path)
            content=hashlib.sha256(canonical({"fields":raw,"geometry":geometry}).encode()).hexdigest()
            identity=hashlib.sha256((str(binding.epoch)+"\n"+locator+"\n"+content).encode()).hexdigest()
            change=SourceChange(change_id=uuid5(binding.epoch,identity), key=binding.source_key.model_copy(update={"canonical_key":key}),
                operation="upsert", position=SourcePosition(stream_id=str(binding.source_key.source_id), epoch=binding.epoch,
                    ordinal=ordinal, source_version=binding.blob.sha256,cursor={"locator":locator}),
                **temporal,received_at=binding.received_at,snapshot=binding.blob,
                content_sha256=content,idempotency_key=identity,raw_fields=raw,geometry=geometry,
                data_nature=options.data_nature,locator=locator)
            output_size+=len(change.model_dump_json().encode())
            if output_size>limits.max_output_bytes-1024: raise InvalidFile("OUTPUT_LIMIT",locator)
            changes.append(change)
        return ParsedFile(changes=tuple(changes),issues=(),complete=True)
    except InvalidFile as error:
        return ParsedFile(changes=tuple(changes[:100]),issues=(issue(error.code,error.locator,error.field),),complete=False)


class FileParser:
    def __init__(self, store, binding: UploadBinding, *, limits=None):
        self.store,self.binding,self.limits=store,binding,limits or ParserLimits()

    def parse(self, blob, key_template, options):
        from lzo.errors import ApiError
        from lzo.semantic.query_process import run_child
        expected="file-v2" if options.temporal_fields is not None else "file-v1"
        if self.binding.parser_contract != expected:
            return incomplete("UNSUPPORTED_PARSER_CONTRACT")
        if blob!=self.binding.blob or key_template!=self.binding.source_key or options!=self.binding.options:
            return incomplete("UPLOAD_BINDING_MISMATCH")
        if blob.size_bytes>self.limits.max_bytes: return incomplete("BYTE_LIMIT")
        media={"csv":{"text/csv","application/csv"},"xlsx":{"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"},
            "json":{"application/json"},"geojson":{"application/geo+json","application/json"}}
        if blob.media_type not in media[options.format]: return incomplete("FORMAT_MISMATCH")
        try:
            with self.store.open(blob) as stream:
                data=stream.read(self.limits.max_bytes+1)
            if len(data)>self.limits.max_bytes: return incomplete("BYTE_LIMIT")
            if len(data)!=blob.size_bytes or hashlib.sha256(data).hexdigest()!=blob.sha256: return incomplete("BLOB_MISMATCH")
            result=run_child({"data":base64.b64encode(data).decode(),"binding":self.binding.model_dump(mode="json"),"limits":asdict(self.limits)},
                module="lzo.ingestion.files", timeout_ms=self.limits.timeout_ms,memory_mb=self.limits.memory_mb,
                input_bytes=self.limits.max_bytes*2+32768,output_bytes=self.limits.max_output_bytes)
            return ParsedFile.model_validate(result)
        except ApiError as error:
            return incomplete("FILE_"+error.code.removeprefix("SPARQL_").removeprefix("FILE_"))
        except Exception:
            return incomplete("FILE_EXECUTION_INCOMPLETE")


if __name__ == "__main__":
    try:
        payload=json.load(sys.stdin)
        limits=ParserLimits(**payload["limits"])
        data=base64.b64decode(payload["data"],validate=True)
        result=decode(data,UploadBinding.model_validate(payload["binding"]),limits)
        output=result.model_dump_json()
        if len(output.encode())>limits.max_output_bytes:
            output=incomplete("OUTPUT_LIMIT").model_dump_json()
        sys.stdout.write(output)
    except (MemoryError,RecursionError):
        sys.stdout.write('{"error":"FILE_RESOURCE_LIMIT"}')
    except Exception:
        sys.stdout.write('{"error":"FILE_EXECUTION_INCOMPLETE"}')
