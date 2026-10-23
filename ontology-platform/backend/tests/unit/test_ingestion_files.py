"""Real bytes exercise normalization, deterministic identity and bounded failures."""
import io
import json
from datetime import datetime, timezone
from types import SimpleNamespace
from uuid import uuid4

import pytest

from lzo.ingestion.files import FileParser, ParserLimits
from lzo.ingestion.schemas import FileOptions, SourceKey, UploadBinding
from lzo.artifacts.store import ArtifactStore


@pytest.fixture
def file_case(tmp_path):
    from openpyxl import Workbook
    import os
    if os.environ.get('LZO_TEST_BLOB_PROVIDER') == 's3':
        from lzo.config import Settings
        from lzo.storage.blobs import get_blob_store
        store = get_blob_store(Settings().model_copy(update={'artifact_max_upload_bytes':2**20,
            'blob_staging_path':tmp_path/'staging'}))
    else:
        store = ArtifactStore(tmp_path, max_upload_bytes=2**20)
    key = SourceKey(source_id=uuid4(), source_system="farm", source_instance="one", collection="pond", canonical_key="")
    options = {fmt: FileOptions(format=fmt, encoding="utf-8", delimiter=",", sheet="pond", header_row=1,
        records_path="", key_paths=("id",), source_crs=None, source_timezone="UTC", data_nature="synthetic")
        for fmt in ("csv", "xlsx", "json", "geojson")}
    book = Workbook(); book.active.title = "pond"
    book.active.append(["id", "name"]); book.active.append(["POND-001", "合成水塘"])
    output = io.BytesIO(); book.save(output)
    data = {"csv": "id,name\nPOND-001,合成水塘\n".encode(), "xlsx": output.getvalue(),
        "json": json.dumps([{"id": "POND-001", "name": "合成水塘"}], ensure_ascii=False).encode(),
        "geojson": json.dumps({"type": "FeatureCollection", "features": [{"type": "Feature",
            "properties": {"id": "POND-001", "name": "合成水塘"},
            "geometry": {"type": "Point", "coordinates": [114.1, 30.2]}}]}).encode()}
    media = {"csv":"text/csv", "xlsx":"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "json":"application/json", "geojson":"application/geo+json"}
    def parse(payload, fmt="json", limits=None, option=None):
        blob = store.put(io.BytesIO(payload), media[fmt])
        opt = option or options[fmt]
        binding = UploadBinding(id=uuid4(), source_key=key, blob=blob, options=opt,
            epoch=uuid4(), received_at=datetime(2026,1,1,tzinfo=timezone.utc), source_revision=1)
        parser = FileParser(store, binding, limits=limits or ParserLimits())
        return parser.parse(blob, key, opt), parser, blob
    return SimpleNamespace(parse=parse, data=data, options=options, key=key, store=store)


@pytest.mark.parametrize("fmt", ["csv", "xlsx", "json", "geojson"])
def test_all_files_emit_same_source_contract(file_case, fmt):
    s = file_case
    result, parser, blob = s.parse(s.data[fmt], fmt)
    assert result.complete and not result.issues
    assert [c.key.canonical_key for c in result.changes] == ['["POND-001"]']
    assert result.changes[0].raw_fields["name"] == "合成水塘"
    assert result.changes[0].data_nature == "synthetic"
    assert parser.parse(blob, s.key, s.options[fmt]) == result
    if fmt == "geojson":
        assert result.changes[0].geometry["coordinates"] == [114.1, 30.2]


@pytest.mark.parametrize("payload,fmt,code", [
    (b'id,name\n1,\xff', "csv", "INVALID_ENCODING"),
    (b'id,id\n1,2', "csv", "DUPLICATE_HEADER"),
    (b'id,name\n1,one,extra', "csv", "ROW_WIDTH"),
    (b'[{"id":"1","id":"2"}]', "json", "DUPLICATE_JSON_KEY"),
    (b'[{"id":"1","value":NaN}]', "json", "NONFINITE_NUMBER"),
    (b'[{"id":"1"},{"id":"1"}]', "json", "DUPLICATE_KEY"),
    (b'[{"id":null}]', "json", "INVALID_KEY"),
    (b'not a workbook', "xlsx", "INVALID_XLSX"),
])
def test_invalid_or_oversized_files_fail_locally(file_case, payload, fmt, code):
    result, _, _ = file_case.parse(payload, fmt)
    assert not result.complete
    assert code in {i.code for i in result.issues}
    assert all(i.locator and "secret" not in i.message for i in result.issues)


def test_formulas_never_become_cached_facts(file_case):
    from openpyxl import Workbook
    book = Workbook(); book.active.title="pond"
    book.active.append(["id","name"]); book.active.append(["1", "=1+1"])
    out=io.BytesIO(); book.save(out)
    parsed, _, _=file_case.parse(out.getvalue(), "xlsx")
    assert not parsed.complete
    assert [(i.code, i.locator) for i in parsed.issues] == [("FORMULA_CELL", "sheet:pond!B2")]


@pytest.mark.parametrize("limits,payload,code", [
    ({"max_bytes": 8}, b'[{"id":"one"}]', "BYTE_LIMIT"),
    ({"max_rows": 1}, b'[{"id":"one"},{"id":"two"}]', "ROW_LIMIT"),
    ({"max_depth": 2}, b'[{"id":"one","v":{"x":{"y":1}}}]', "DEPTH_LIMIT"),
    ({"max_text": 4}, b'[{"id":"private-secret"}]', "TEXT_LIMIT"),
])
def test_resource_budgets_are_independent(file_case, limits, payload, code):
    parsed, _, _=file_case.parse(payload, limits=ParserLimits(**limits))
    assert not parsed.complete and code in {i.code for i in parsed.issues}


def test_actual_zip_expansion_is_bounded(file_case):
    parsed, _, _=file_case.parse(file_case.data["xlsx"], "xlsx", ParserLimits(max_expanded_bytes=256))
    assert not parsed.complete and "ZIP_EXPANSION_LIMIT" in {i.code for i in parsed.issues}


def test_unregistered_options_and_keys_cannot_change_authority(file_case):
    s=file_case
    _, parser, blob=s.parse(s.data["json"])
    result=parser.parse(blob, s.key, s.options["json"].model_copy(update={"data_nature":"verified"}))
    assert not result.complete and result.issues[0].code == "UPLOAD_BINDING_MISMATCH"


def test_unsupported_registered_parser_contract_cannot_be_silently_redecoded(file_case):
    s=file_case
    _,parser,blob=s.parse(s.data['json'])
    unsupported=parser.binding.model_copy(update={'parser_contract':'future-incompatible'})
    result=FileParser(s.store,unsupported).parse(blob,s.key,s.options['json'])
    assert not result.complete and result.issues[0].code=='UNSUPPORTED_PARSER_CONTRACT'


@pytest.mark.parametrize("limits", [{"timeout_ms":1},{"memory_mb":1},{"max_output_bytes":100}])
def test_actual_child_failures_never_return_complete(file_case,limits):
    result,_,_=file_case.parse(file_case.data["json"],limits=ParserLimits(**limits))
    assert not result.complete and result.issues


def test_supervisor_refuses_unregistered_executable_selector():
    from lzo.semantic.query_process import run_child
    with pytest.raises(ValueError,match="Unregistered"):
        run_child({},module="os",timeout_ms=1000,memory_mb=128,input_bytes=1024,output_bytes=1024)


@pytest.mark.parametrize("geometry", [
    {"type":"Point","coordinates":[[1,2]]},
    {"type":"LineString","coordinates":[[1,2]]},
    {"type":"GeometryCollection","geometries":[{"type":"Unknown","coordinates":[1,2]}]},
])
def test_geojson_geometry_shape_is_validated_without_coordinate_guessing(file_case,geometry):
    raw=json.dumps({"type":"FeatureCollection","features":[{"type":"Feature","properties":{"id":"1"},"geometry":geometry}]}).encode()
    result,_,_=file_case.parse(raw,"geojson")
    assert not result.complete and result.issues[0].code=="INVALID_GEOMETRY"


@pytest.mark.parametrize("mode,code",[("entity","INVALID_XLSX"),("external","EXTERNAL_OR_MACRO_REFERENCE"),("dimension","DIMENSION_LIMIT")])
def test_xlsx_rejects_entities_external_references_and_sparse_dimensions(file_case,mode,code):
    import zipfile
    output=io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(file_case.data["xlsx"])) as source,zipfile.ZipFile(output,"w",zipfile.ZIP_DEFLATED) as target:
        for name in source.namelist():
            data=source.read(name)
            if mode=="entity" and name=="xl/worksheets/sheet1.xml":
                data=b'<!DOCTYPE x [<!ENTITY x "private-secret">]>'+data.replace(b"POND-001",b"&x;")
            if mode=="external" and name=="xl/_rels/workbook.xml.rels":
                data=data.replace(b"<Relationship ",b'<Relationship TargetMode="External" ',1)
            if mode=="dimension" and name=="xl/worksheets/sheet1.xml":
                data=data.replace(b'A1:B2',b'A1:XFD1048576')
            target.writestr(name,data)
    result,_,_=file_case.parse(output.getvalue(),"xlsx")
    assert not result.complete and result.issues[0].code==code


def workbook_with_rows(rows):
    from openpyxl import Workbook
    book=Workbook(); book.active.title="pond"
    for row in rows: book.active.append(row)
    output=io.BytesIO(); book.save(output); book.close()
    return output.getvalue()


def replace_worksheet_rows(payload, rows):
    import re
    import zipfile
    output=io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(payload)) as source, zipfile.ZipFile(output,"w",zipfile.ZIP_DEFLATED) as target:
        for name in source.namelist():
            data=source.read(name)
            if name=="xl/worksheets/sheet1.xml":
                data=re.sub(b'<sheetData>.*?</sheetData>', b'<sheetData>'+rows+b'</sheetData>',data)
            target.writestr(name,data)
    return output.getvalue()


@pytest.mark.parametrize("rows,code,location", [
    (b'<row r="2"><c r="A2" t="n"><v>1</v></c><c r="B2"><f>1+1</f><v>2</v></c><c r="B2" t="n"><v>9</v></c></row>', "FORMULA_CELL", "B2"),
    (b'<row r="2"><c r="A2" t="n"><v>1</v></c><c r="A2" t="n"><v>2</v></c></row>', "INVALID_COORDINATE", "A2"),
    (b'<row r="2"><c r="B2" t="n"><v>1</v></c><c r="A2" t="n"><v>2</v></c></row>', "INVALID_COORDINATE", "A2"),
    (b'<row r="2"><c r="A2" t="n"><v>1</v></c></row><row r="2"><c r="A2" t="n"><v>2</v></c></row>', "INVALID_COORDINATE", "row:2"),
    (b'<row r="3"><c r="A3" t="n"><v>1</v></c></row><row r="2"><c r="A2"><f>1+1</f><v>2</v></c></row>', "INVALID_COORDINATE", "row:2"),
    (b'<row r="2"><c r="A3" t="n"><v>1</v></c></row>', "INVALID_COORDINATE", "A3"),
    (b'<row r="2"><c t="n"><v>1</v></c></row>', "INVALID_COORDINATE", "row:2"),
    (b'<row><c r="A2" t="n"><v>1</v></c></row>', "INVALID_COORDINATE", "row:unknown"),
    (b'<row r="2"><c r="XFD2" t="n"><v>1</v></c></row>', "DIMENSION_LIMIT", "XFD2"),
    (b'<row r="1048576"><c r="A1048576" t="n"><v>1</v></c></row>', "DIMENSION_LIMIT", "row:1048576"),
])
def test_xlsx_actual_coordinates_cannot_hide_or_overwrite_cells(file_case,rows,code,location):
    header=b'<row r="1"><c r="A1" t="inlineStr"><is><t>id</t></is></c><c r="B1" t="inlineStr"><is><t>name</t></is></c></row>'
    payload=replace_worksheet_rows(file_case.data["xlsx"],header+rows)
    result,_,_=file_case.parse(payload,"xlsx")
    assert not result.complete
    assert [(i.code,i.locator) for i in result.issues]==[(code,"sheet:pond!"+location)]


@pytest.mark.parametrize("row,expected", [
    (["one",None,None], {"id":"one","middle":None,"last":None}),
    (["one",None,"last"], {"id":"one","middle":None,"last":"last"}),
    (["one","middle",None], {"id":"one","middle":"middle","last":None}),
])
def test_xlsx_sparse_fields_preserve_header_width_and_row_locator(file_case,row,expected):
    payload=workbook_with_rows([["id","middle","last"],row])
    result,_,_=file_case.parse(payload,"xlsx")
    assert result.complete and not result.issues
    assert len(result.changes)==1
    assert result.changes[0].raw_fields==expected
    assert result.changes[0].locator=="sheet:pond!row:2"


@pytest.mark.parametrize("row,code,location", [
    ([None,"present"],"INVALID_KEY","row:2"),
    (["one",None,"extra"],"ROW_WIDTH","row:2"),
    (["one",None,"=1+1"],"FORMULA_CELL","C2"),
])
def test_xlsx_sparse_padding_cannot_hide_invalid_keys_or_excess_cells(file_case,row,code,location):
    result,_,_=file_case.parse(workbook_with_rows([["id","optional"],row]),"xlsx")
    assert not result.complete
    assert [(i.code,i.locator) for i in result.issues]==[(code,"sheet:pond!"+location)]


@pytest.mark.parametrize("nested", [False,True])
@pytest.mark.parametrize("valid", [False,True])
def test_geojson_multiline_components_obey_line_cardinality_at_every_depth(file_case,nested,valid):
    geometry={"type":"MultiLineString","coordinates":[[[1,2],[3,4]],[[5,6],[7,8]] if valid else [[5,6]]]}
    if nested:
        geometry={"type":"GeometryCollection","geometries":[{"type":"GeometryCollection","geometries":[geometry]}]}
    payload=json.dumps({"type":"FeatureCollection","crs":{"type":"name","properties":{"name":"EPSG:4326"}},
        "features":[{"type":"Feature","properties":{"id":"one"},"geometry":geometry}]}).encode()
    result,_,_=file_case.parse(payload,"geojson")
    if valid:
        assert result.complete and len(result.changes)==1
        assert result.changes[0].geometry=={**geometry,"crs":{"type":"name","properties":{"name":"EPSG:4326"}}}
    else:
        assert not result.complete
        assert [(i.code,i.locator) for i in result.issues]==[("INVALID_GEOMETRY","$/0")]
